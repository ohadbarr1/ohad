"""Format-drift tests for the asset-list normaliser: python3 -m unittest pipeline/asset_lists/test_normalise.py

Each case is a small workbook in the CMA layout with one of the variants seen across managers (23.10.26):
  header row below title rows (Meitav, Migdal), the 'לא סחיר אופציות' header "שווי הוגן (בש"ח)" (all managers before Q1'26),
  percentages as fractions vs percent, the "ריק במקור" filler (Meitav, Analyst), IL ISINs vs TASE numbers (Altshuler vs Meitav),
  the filer's summary sheet out of step on one sheet (Altshuler bonds, 4 files), helper sheets to ignore (Migdal 'File Name Info').
"""
from __future__ import annotations
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from asset_lists import normalise as N  # noqa: E402
from asset_lists.common import last_quarters, parse_name, qlabel  # noqa: E402

FV = 'שווי הוגן (באלפי ש"ח)'
HDR = ["מספר קופה/קרן/ח.פ. עבור חברת ביטוח", "מספר מסלול", "שם מנפיק", "מספר מנפיק", "סוג מספר מזהה מנפיק", "שם נייר ערך", "מספר נייר ערך",
       'ישראל/חו"ל', "מדינה לפי חשיפה כלכלית", "זירת מסחר", "מטבע פעילות", FV, "שיעור מנכסי אפיק ההשקעה", "שיעור מסך נכסי ההשקעה"]


def book(sheets: dict, summary: dict | None = None, title_rows: int = 0) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "עמוד פתיחה"
    for r in (["יש לבחור תחום:", "נכסי עמיתים - קופות גמל"], ["יש לבחור את רבעון הדיווח:", 2], ["יש לבחור את שנת הדיווח:", 2026],
              ["יש לבחור את הגוף המוסדי:", "גוף בדיקה בע\"מ"], ["ח.פ. הגוף המוסדי:", 512345678]):
        ws.append(r)
    if summary is not None:
        s = wb.create_sheet("סכום נכסים")
        s.append(["ידווח בקבצי נכסי הנוסטרו בלבד"])
        s.append(["אפיק", FV])
        for k, v in summary.items():
            s.append([k, v])
    for name, (hdr, rows) in sheets.items():
        s = wb.create_sheet(name)
        for _ in range(title_rows):
            s.append(["כותרת"])
        s.append(hdr)
        for r in rows:
            s.append(r)
    p = Path(tempfile.mkdtemp()) / "t.xlsx"
    wb.save(p)
    return p


def eq_row(track, issuer, iid, sec, fv, pct, country="ישראל", venue="TASE"):
    return [track, track, issuer, iid, "ח.פ.", "x", sec, "ישראל", country, venue, "ILS", fv, None, pct]


class TestNormalise(unittest.TestCase):
    def test_basic_and_isin(self):
        p = book({"מניות מבכ ויהש": (HDR, [eq_row(1, "לאומי", 520018078, "IL0006046119", 60.0, 0.6), eq_row(1, "פועלים", 520000118, "662577", 40.0, 0.4)])},
                 summary={"מניות מבכ ויהש": 100.0})
        rows, info = N.read_file(p)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["sec_id"], "604611")         # IL ISIN -> TASE number
        self.assertEqual(rows[1]["sec_id"], "662577")
        self.assertEqual(info["cover"]["year"], "2026")
        v = N.validate(rows, info)
        self.assertTrue(v["ok"])
        self.assertEqual(v["warnings"], 0)

    def test_title_rows_percent_and_placeholder(self):
        rows_in = [eq_row(7, "ריק במקור", "ריק במקור", "1234567", 30.0, 30.0, country="ריק במקור"),
                   eq_row(7, "ב", 520000001, "7654321", 70.0, 70.0)]
        p = book({"מניות מבכ ויהש": (HDR, rows_in)}, title_rows=3)
        rows, info = N.read_file(p)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["issuer"], "")                # "ריק במקור" is a blank
        self.assertEqual(rows[0]["country"], "")
        self.assertTrue(N.validate(rows, info)["ok"])          # 30 + 70 read as percent

    def test_options_header_alias(self):
        hdr = [h if h != FV else 'שווי הוגן (בש"ח)' for h in HDR]
        p = book({"לא סחיר אופציות": (hdr, [eq_row(1, "מזרחי", 520000522, "406259820", 381.8, 1.0)])}, summary={"לא סחיר אופציות": 381.8})
        rows, info = N.read_file(p)
        self.assertEqual(len(rows), 1)
        self.assertAlmostEqual(info["sheets"]["לא סחיר אופציות"]["total"], 381.8)
        self.assertEqual(N.validate(rows, info)["summary"]["fail"], [])

    def test_summary_out_of_step_is_warning(self):
        p = book({"מניות מבכ ויהש": (HDR, [eq_row(1, "א", 520000001, "111111", 1000.0, 0.5)]),
                  "איגרות חוב": (HDR, [eq_row(1, "ב", 520000002, "222222", 1000.0, 0.5)])},
                 summary={"מניות מבכ ויהש": 1000.0, "איגרות חוב": 1006.0})
        rows, info = N.read_file(p)
        v = N.validate(rows, info)
        self.assertEqual(len(v["summary"]["fail"]), 1)         # the bonds sheet misses by 0.6%
        self.assertTrue(v["ok"])                               # the grand total ties within 0.5%

    def test_unknown_sheets_and_off_balance(self):
        commit_hdr = ["מספר קופה/קרן/ח.פ. עבור חברת ביטוח", "מספר מסלול", "שם קרן השקעה", 'יתרת המחויבות לתקופת הדיווח (באלפי ש"ח)']
        p = book({"מניות מבכ ויהש": (HDR, [eq_row(1, "א", 520000001, "111111", 10.0, 1.0)]),
                  "יתרות התחייבות להשקעה": (commit_hdr, [[1, 1, "קרן", 55.0], [1, 1, "קרן ב", 45.0]]),
                  "File Name Info": (["x"], [[1]])})
        rows, info = N.read_file(p)
        self.assertEqual(info["off_balance"]["commitments"], {"1": 100.0})
        self.assertIn("File Name Info", info["unknown_sheets"])
        self.assertEqual(len(rows), 1)

    def test_names(self):
        self.assertEqual(parse_name("/media/x/513173393_gm_p_0226.xlsx"), ("513173393", "gm", "2026Q2"))
        self.assertEqual(parse_name("513621110_GM_P_0125.xlsx"), ("513621110", "gm", "2025Q1"))
        self.assertEqual(parse_name("513026484_pn_p_0226_for-yesod.xlsx"), ("513026484", "pn", "2026Q2"))
        self.assertIsNone(parse_name("511880460_g811_P423.xlsx"))  # per-track files of the old format are not used
        self.assertEqual(last_quarters(3, "2026Q1"), ["2025Q3", "2025Q4", "2026Q1"])
        self.assertEqual(qlabel("2026Q2"), "Q2'26")


if __name__ == "__main__":
    unittest.main()
